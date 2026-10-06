"""Extraction must report segments, honest quality, and explicit warnings."""

import io
import unittest

import _support  # noqa: F401
from extractors import (
    decode_text_bytes,
    extract_from_image,
    extract_from_pdf,
    extract_from_pptx,
    extract_from_text,
)

LONG_TEXT = (
    "Operating systems coordinate hardware and software resources. "
    "A process is a program in execution, while a thread is a unit of execution "
    "within a process. The scheduler decides which process runs next using "
    "algorithms such as round robin, shortest job first, and priority scheduling. "
    "Virtual memory lets processes use more memory than is physically available "
    "by moving inactive pages to disk. Deadlock occurs when processes wait for "
    "each other indefinitely, and it can be prevented by breaking one of the "
    "necessary conditions. File systems organize data on storage devices using "
    "inodes, directories, and allocation tables. "
) * 3


class TextExtractionTests(unittest.TestCase):
    def test_long_text_is_good_quality_with_one_segment(self):
        result = extract_from_text(LONG_TEXT, "notes.txt")
        self.assertEqual(result["file_type"], "text")
        self.assertEqual(result["quality"], "good")
        self.assertEqual(result["warnings"], [])
        self.assertEqual(len(result["segments"]), 1)
        self.assertEqual(result["segments"][0]["label"], "Text")
        self.assertEqual(result["segments"][0]["source_type"], "text")
        self.assertIn("words", result["metadata"])

    def test_short_text_warns_about_limits(self):
        result = extract_from_text("Too short to review.", "tiny.txt")
        self.assertEqual(result["quality"], "partial")
        self.assertTrue(any("very short" in w for w in result["warnings"]))

    def test_empty_text_is_poor(self):
        result = extract_from_text("   ", "empty.txt")
        self.assertEqual(result["quality"], "poor")
        self.assertEqual(result["warnings"], ["No text was provided."])

    def test_markdown_line_endings_preserved(self):
        result = extract_from_text("line one\r\nline two\r\n", "md.txt")
        self.assertIn("line one", result["text"])
        self.assertIn("line two", result["text"])
        self.assertNotIn("\r", result["text"])


class PdfExtractionTests(unittest.TestCase):
    def _make_pdf(self, pages):
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        buffer = io.BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=letter)
        for text in pages:
            if text:
                pdf.drawString(72, 720, text)
            pdf.showPage()
        pdf.save()
        buffer.seek(0)
        return buffer

    def test_blank_pdf_reports_no_ocr_honestly(self):
        result = extract_from_pdf(self._make_pdf(["", ""]))
        self.assertEqual(result["file_type"], "pdf")
        self.assertEqual(result["page_count"], 2)
        self.assertEqual(result["quality"], "poor")
        self.assertTrue(any("No OCR was performed" in w for w in result["warnings"]))
        self.assertEqual(result["text"], "")

    def test_partially_scanned_pdf_warns_about_missing_pages(self):
        result = extract_from_pdf(self._make_pdf(["", LONG_TEXT[:400]]))
        self.assertEqual(result["page_count"], 2)
        self.assertEqual(result["quality"], "partial")
        self.assertTrue(
            any("1 of 2 pages/slides" in w for w in result["warnings"]),
            result["warnings"],
        )

    def test_text_pdf_segments_and_markers(self):
        result = extract_from_pdf(self._make_pdf([LONG_TEXT[:500], LONG_TEXT[:500]]))
        self.assertEqual(result["quality"], "good")
        self.assertEqual(len(result["segments"]), 2)
        self.assertEqual(result["segments"][0]["source_type"], "pdf_page")
        self.assertEqual(result["segments"][1]["label"], "Page 2")
        self.assertIn("--- [Page 1] ---", result["text"])
        self.assertIn("--- [Page 2] ---", result["text"])


class PptxExtractionTests(unittest.TestCase):
    def _make_pptx(self):
        from pptx import Presentation

        prs = Presentation()
        filler = LONG_TEXT[:420]

        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Operating Systems Overview"
        slide.placeholders[1].text = filler

        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Memory Management"
        slide2.placeholders[1].text = (
            "- Paging divides memory into fixed size frames.\n"
            "- Segmentation divides memory into variable size segments.\n"
            f"- {filler}"
        )

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer

    def test_slides_become_segments_without_title_duplication(self):
        result = extract_from_pptx(self._make_pptx())
        self.assertEqual(result["file_type"], "pptx")
        self.assertEqual(result["page_count"], 2)
        self.assertEqual(len(result["segments"]), 2)
        self.assertEqual(result["segments"][0]["label"], "Slide 1")
        self.assertEqual(result["segments"][1]["source_type"], "pptx_slide")
        self.assertIn("--- [Slide 1] ---", result["text"])
        self.assertIn("--- [Slide 2] ---", result["text"])

        first_slide_text = result["segments"][0]["text"]
        self.assertTrue(first_slide_text.startswith("Title: Operating Systems Overview"))
        self.assertEqual(first_slide_text.count("Operating Systems Overview"), 1)
        self.assertEqual(first_slide_text.count(filler_marker()), 1)

    def test_slides_with_enough_text_are_good_quality(self):
        result = extract_from_pptx(self._make_pptx())
        self.assertEqual(result["quality"], "good", result["warnings"])
        self.assertEqual(result["warnings"], [])


def filler_marker():
    return LONG_TEXT[:420]


class ImageExtractionTests(unittest.TestCase):
    def test_image_reports_no_ocr_and_returns_base64(self):
        from PIL import Image

        buffer = io.BytesIO()
        Image.new("RGB", (32, 32), color=(10, 20, 200)).save(buffer, format="PNG")
        buffer.seek(0)

        result = extract_from_image(buffer, "diagram.png")
        self.assertEqual(result["file_type"], "image")
        self.assertEqual(result["mime_type"], "image/png")
        self.assertTrue(result["image_b64"])
        self.assertEqual(result["quality"], "partial")
        self.assertTrue(any("No OCR was performed" in w for w in result["warnings"]))
        self.assertIn("No text was extracted from this image", result["text"])
        self.assertEqual(result["segments"][0]["source_type"], "image")


class DecodeTests(unittest.TestCase):
    def test_utf16_bom_decoded(self):
        self.assertEqual(decode_text_bytes("Café".encode("utf-16"), "a.txt"), "Café")

    def test_utf8_bom_decoded(self):
        self.assertEqual(
            decode_text_bytes(b"\xef\xbb\xbfhello", "a.txt"), "hello"
        )

    def test_undecodable_bytes_raise(self):
        with self.assertRaises(ValueError):
            decode_text_bytes(b"\x81\x81", "binary.bin")


if __name__ == "__main__":
    unittest.main()
