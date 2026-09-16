import os
import io
import base64
import logging
from typing import Dict, Any, Optional

logger = logging.getLogger(__name__)

def extract_from_pdf(file_stream: io.BytesIO) -> Dict[str, Any]:
    """Extract text and metadata from PDF stream using pypdf."""
    try:
        import pypdf
        reader = pypdf.PdfReader(file_stream)
        pages_text = []
        total_pages = len(reader.pages)
        
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            cleaned = page_text.strip()
            if cleaned:
                pages_text.append(f"--- [Page {idx + 1}] ---\n{cleaned}")
                
        full_text = "\n\n".join(pages_text)
        return {
            "text": full_text,
            "page_count": total_pages,
            "file_type": "pdf",
            "metadata": {"pages": total_pages, "characters": len(full_text)}
        }
    except Exception as e:
        logger.error(f"Error reading PDF: {e}")
        raise ValueError(f"Failed to read PDF file: {str(e)}")

def extract_from_pptx(file_stream: io.BytesIO) -> Dict[str, Any]:
    """Extract slide titles, bullet points, tables, and notes from PPTX."""
    try:
        from pptx import Presentation
        prs = Presentation(file_stream)
        slides_text = []
        total_slides = len(prs.slides)
        
        for idx, slide in enumerate(prs.slides):
            slide_parts = [f"--- [Slide {idx + 1}] ---"]
            
            # Slide Title
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_parts.append(f"Title: {slide.shapes.title.text.strip()}")
            
            # Shapes & Text Frames
            for shape in slide.shapes:
                if shape.has_text_frame and shape != slide.shapes.title:
                    frame_text = shape.text_frame.text.strip()
                    if frame_text:
                        slide_parts.append(frame_text)
                elif shape.has_table:
                    # Extract table content
                    table_rows = []
                    for row in shape.table.rows:
                        row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                        table_rows.append(" | ".join(row_vals))
                    if table_rows:
                        slide_parts.append("Table:\n" + "\n".join(table_rows))
            
            # Speaker notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_parts.append(f"Notes: {notes}")
            
            content = "\n".join(slide_parts)
            slides_text.append(content)
            
        full_text = "\n\n".join(slides_text)
        return {
            "text": full_text,
            "page_count": total_slides,
            "file_type": "pptx",
            "metadata": {"slides": total_slides, "characters": len(full_text)}
        }
    except Exception as e:
        logger.error(f"Error reading PPTX: {e}")
        raise ValueError(f"Failed to read PowerPoint file: {str(e)}")

def extract_from_image(file_stream: io.BytesIO, filename: str) -> Dict[str, Any]:
    """Analyze image file, extract dimensions and encode as base64 for multimodal AI."""
    try:
        from PIL import Image
        file_bytes = file_stream.read()
        img = Image.open(io.BytesIO(file_bytes))
        width, height = img.size
        img_format = (img.format or "JPEG").lower()
        
        # Determine mime type
        mime_type = "image/jpeg"
        if img_format in ["png"]:
            mime_type = "image/png"
        elif img_format in ["webp"]:
            mime_type = "image/webp"
            
        b64_data = base64.b64encode(file_bytes).decode('utf-8')
        
        # Summary description for fallback/metadata
        description = f"Lesson image ({width}x{height}px, {img_format.upper()}) containing lesson material, diagrams, or notes."
        
        return {
            "text": f"[Uploaded Image: {filename} ({width}x{height})]\nVisual learning material containing diagrams, equations, or notes.",
            "page_count": 1,
            "file_type": "image",
            "image_b64": b64_data,
            "mime_type": mime_type,
            "metadata": {
                "width": width,
                "height": height,
                "format": img_format,
                "filename": filename
            }
        }
    except Exception as e:
        logger.error(f"Error reading image: {e}")
        raise ValueError(f"Failed to read image file: {str(e)}")

def extract_from_text(text: str, filename: str = "Pasted Text") -> Dict[str, Any]:
    """Format and normalize plain text or markdown lesson input."""
    clean_text = text.strip()
    words = len(clean_text.split())
    # Estimate pages (approx 400 words per page)
    est_pages = max(1, round(words / 400))
    return {
        "text": clean_text,
        "page_count": est_pages,
        "file_type": "text",
        "metadata": {"words": words, "characters": len(clean_text), "filename": filename}
    }

def process_file_input(file_obj, filename: str) -> Dict[str, Any]:
    """Unified file processor based on extension."""
    ext = os.path.splitext(filename)[1].lower()
    stream = io.BytesIO(file_obj.read())
    
    if ext == ".pdf":
        return extract_from_pdf(stream)
    elif ext in [".pptx", ".ppt"]:
        return extract_from_pptx(stream)
    elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp"]:
        return extract_from_image(stream, filename)
    elif ext in [".txt", ".md", ".csv"]:
        stream.seek(0)
        content = stream.read().decode("utf-8", errors="ignore")
        return extract_from_text(content, filename)
    else:
        # Try reading as text
        stream.seek(0)
        try:
            content = stream.read().decode("utf-8")
            return extract_from_text(content, filename)
        except Exception:
            raise ValueError(f"Unsupported file format: {ext}. Please upload a PDF, PPTX, image, or text file.")
